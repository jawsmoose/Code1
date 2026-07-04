#!/usr/bin/env python3
"""
Build the "Project Review" Power BI project (PBIP) + its CSV data files.

Replicates the one-page "Transmission and Grid — Project Review" layout:
  * Top half   : project attribute cards, fed from data/projects.csv
  * Bottom half: cross-workstream milestone timeline (scatter chart) whose
                 circles move via Milestones[TimelinePos] and recolor via
                 Milestones[Status]  (On Track / At Risk / Delayed / Not Started)

Pages emitted: Template (P00), five project pages (P01–P05), Portfolio Summary.

Run:  python3 tools/build_pbip.py   (from the ProjectReview/ folder)
Everything is regenerated deterministically; edit data in the CSVs, not here,
unless you want to change the layout itself.
"""

import csv
import json
import os
import sys

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DATA_DIR = os.path.join(ROOT, "data")
REPORT_DIR = os.path.join(ROOT, "Project Review.Report")
MODEL_DIR = os.path.join(ROOT, "Project Review.SemanticModel")

PAGE_W, PAGE_H = 1280, 720

# ----- status palette (validated: CVD-separated, >=3:1 on white) -----
STATUS_COLORS = {
    "On Track": "#4E8542",
    "At Risk": "#C4680C",
    "Delayed": "#B3261E",
    "Not Started": "#787E88",
}
NAVY = "#1F3252"        # section header bars
INK = "#252423"         # primary text
MUTED = "#5F6368"       # card titles
CARD_BORDER = "#C9CDD4"
FONT = "Segoe UI"

# =====================================================================
# 1. DATA
# =====================================================================

FIELDS = [
    # (csv column, measure name, card title, template placeholder)
    ("ReviewDate", "v Review Date", "REVIEW DATE", "[ MM / DD / YYYY ]"),
    ("ProjectName", "v Project Name", "PROJECT", "[ Project Name ]"),
    ("LeadInterconnection", "v Lead IX", "INTERCONNECTION", "[ Name ]"),
    ("LeadHVEngineering", "v Lead HV Eng", "HV ENGINEERING", "[ Name ]"),
    ("LeadSCADATelecom", "v Lead SCADA", "SCADA / TELECOM", "[ Name ]"),
    ("LeadHVOM", "v Lead HV OM", "HV O&M", "[ Name ]"),
    ("LeadAssetMgmt", "v Lead Asset Mgmt", "ASSET MGMT", "[ Name ]"),
    ("ProjectState", "v Project State", "PROJECT STATE", "[ State ]"),
    ("TOUtility", "v TO Utility", "TO / UTILITY", "[ TO / Utility ]"),
    ("ISORTO", "v ISO RTO", "ISO / RTO", "[ ISO / RTO ]"),
    ("OTBEPAgreement", "v OTB EP", "OTB / E&P AGREEMENT", "[ Yes / No ] (add date(s) if Y)"),
    ("POIVoltage", "v POI Voltage", "POI VOLTAGE", "[ ___ kV ]"),
    ("POIType", "v POI Type", "POI TYPE", "[ New SS / Existing / BTM ]"),
    ("Solar", "v Solar", "SOLAR (MW AC / DC)", "[ ___ AC / ___ DC ]"),
    ("Storage", "v Storage", "STORAGE (MW / MWH)", "[ ___ MW / ___ MWh ]"),
    ("PGIAGIA", "v PGIA GIA", "PGIA / GIA", "[ MM / YYYY ]"),
    ("COD", "v COD", "COD", "[ MM / YYYY ]"),
    ("EPCHV", "v EPC HV", "EPC / HV", "[ Contractor ]"),
    ("OMProvider", "v OM Provider", "O&M PROVIDER", "[ Provider ]"),
    ("POIISD", "v POI ISD", "POI ISD", "[ MM / YYYY ]"),
    ("POIISDStatus", "v POI ISD Status", "POI ISD STATUS", "[ On Track ]"),
    ("EstIXCost", "v Est IX Cost", "EST. IX COST", "[ $ ___ M ]"),
    ("IXCostStatus", "v IX Cost Status", "IX COST STATUS", "[ Within Budget ]"),
    ("SolarInverters", "v Solar Inverters", "SOLAR INVERTERS", "[ Make / Model / Qty ]"),
    ("MPT", "v MPT", "MPT", "[ ___ MVA · Δ-Yg / Yg-Yg ]"),
    ("StoragePCS", "v Storage PCS", "STORAGE / PCS", "[ Make / Model / Qty / Duration ]"),
    ("GenTieLength", "v Gen Tie", "GEN-TIE LENGTH", "[ Miles ]"),
    ("ReactivePowerComp", "v Reactive Power", "REACTIVE POWER COMPENSATION", "[ Type / Size / Qty ]"),
    ("PowerPlantCtrl", "v Power Plant Ctrl", "POWER PLANT CTRL", "[ Vendor ]"),
    ("SafeHarbor", "v Safe Harbor", "SAFE HARBOR", "[ MPT ~ Qty / kV-MVA / FID Year ]"),
    ("OtherAttributes", "v Other", "OTHER", "[ NITS, Fuel / Equip Change, etc. ]"),
    ("FutureExpansion", "v Future Expansion", "FUTURE EXPANSION — NEW APPS", "[ Capacity / Surplus / Load IX - DC ]"),
    ("TransmissionNUs", "v Transmission NUs", "TRANSMISSION SYSTEM NUs", "[ Desc. / ISD ]"),
    ("TopGatingItem", "v Top Gating Item", "TOP GATING ITEM", "[ Top gating item ]"),
]
FIELD = {f[0]: f for f in FIELDS}

PROJECT_COLUMNS = ["ProjectID"] + [f[0] for f in FIELDS]

PROJECTS = [
    dict(ProjectID="P00", **{col: ph for col, _m, _t, ph in FIELDS}),
    dict(
        ProjectID="P01", ReviewDate="07/01/2026", ProjectName="Sunport Ridge Solar + Storage",
        LeadInterconnection="K. Rivera", LeadHVEngineering="A. Chen", LeadSCADATelecom="D. Okafor",
        LeadHVOM="M. Salinas", LeadAssetMgmt="J. Whitfield",
        ProjectState="Texas", TOUtility="Oncor", ISORTO="ERCOT",
        OTBEPAgreement="Yes (E&P 03/2025)", POIVoltage="345 kV", POIType="New SS",
        Solar="200 AC / 260 DC", Storage="100 MW / 400 MWh", PGIAGIA="08/2024", COD="06/2027",
        EPCHV="Blattner", OMProvider="NovaSource",
        POIISD="12/2026", POIISDStatus="On Track", EstIXCost="$14.2 M", IXCostStatus="Within Budget",
        SolarInverters="Sungrow SG4400UD-MV / 46", MPT="230 MVA · Δ-Yg",
        StoragePCS="Tesla Megapack 2XL / 100 / 4h", GenTieLength="3.4 miles",
        ReactivePowerComp="Cap Bank / 40 MVAr / 2", PowerPlantCtrl="Emerson",
        SafeHarbor="MPT ~ 1 / 345kV-230MVA / FID 2024", OtherAttributes="NITS filed",
        FutureExpansion="Surplus IX 40 MW - DC", TransmissionNUs="None",
        TopGatingItem="Finalize PEC scope & schedule with Oncor",
    ),
    dict(
        ProjectID="P02", ReviewDate="07/01/2026", ProjectName="Mesa Verde Solar",
        LeadInterconnection="S. Patel", LeadHVEngineering="R. Dominguez", LeadSCADATelecom="L. Nguyen",
        LeadHVOM="T. Brooks", LeadAssetMgmt="C. Ferraro",
        ProjectState="Arizona", TOUtility="APS", ISORTO="WECC (non-RTO)",
        OTBEPAgreement="No", POIVoltage="230 kV", POIType="Existing",
        Solar="150 AC / 195 DC", Storage="75 MW / 300 MWh", PGIAGIA="02/2025", COD="11/2027",
        EPCHV="Rosendin", OMProvider="SOLV Energy",
        POIISD="05/2027", POIISDStatus="At Risk", EstIXCost="$9.8 M", IXCostStatus="Over (+$1.1 M)",
        SolarInverters="GE FLEXINVERTER / 38", MPT="180 MVA · Yg-Yg",
        StoragePCS="Fluence Gridstack / 75 / 4h", GenTieLength="1.2 miles",
        ReactivePowerComp="None (study pending)", PowerPlantCtrl="Trimark",
        SafeHarbor="None", OtherAttributes="Fuel / equip change — N/A",
        FutureExpansion="Load IX 20 MW - DC", TransmissionNUs="115 kV line rebuild / 04/2027",
        TopGatingItem="IX cost overrun mitigation plan with APS",
    ),
    dict(
        ProjectID="P03", ReviewDate="07/01/2026", ProjectName="Prairie Sky Solar + Storage",
        LeadInterconnection="E. Kowalski", LeadHVEngineering="H. Yamada", LeadSCADATelecom="B. Ansari",
        LeadHVOM="G. Martin", LeadAssetMgmt="N. Adeyemi",
        ProjectState="Illinois", TOUtility="ComEd", ISORTO="PJM",
        OTBEPAgreement="No", POIVoltage="138 kV", POIType="New SS",
        Solar="300 AC / 390 DC", Storage="150 MW / 600 MWh", PGIAGIA="11/2025 (target)", COD="09/2028",
        EPCHV="TBD (RFP Q3 2026)", OMProvider="TBD",
        POIISD="03/2028", POIISDStatus="On Track", EstIXCost="$22.5 M", IXCostStatus="Within Budget",
        SolarInverters="TBD (shortlist: SMA / Sungrow)", MPT="345 MVA · Δ-Yg (LOI)",
        StoragePCS="TBD / 150 / 4h", GenTieLength="6.8 miles",
        ReactivePowerComp="STATCOM / 75 MVAr / 1 (prelim)", PowerPlantCtrl="TBD",
        SafeHarbor="MPT ~ 2 / 138kV-175MVA / FID 2025", OtherAttributes="NITS + equip change pending",
        FutureExpansion="Capacity +100 MW phase 2", TransmissionNUs="Desc. TBD / ISD 2028",
        TopGatingItem="GIA execution with PJM / ComEd",
    ),
    dict(
        ProjectID="P04", ReviewDate="07/01/2026", ProjectName="Cypress Bend Solar",
        LeadInterconnection="V. Osei", LeadHVEngineering="P. Lindqvist", LeadSCADATelecom="F. Cardona",
        LeadHVOM="I. Novak", LeadAssetMgmt="W. Tran",
        ProjectState="Georgia", TOUtility="Georgia Power", ISORTO="SERC (non-RTO)",
        OTBEPAgreement="Yes (04/2026)", POIVoltage="115 kV", POIType="Existing",
        Solar="120 AC / 156 DC", Storage="50 MW / 200 MWh", PGIAGIA="06/2024", COD="12/2026",
        EPCHV="McCarthy", OMProvider="Self-perform",
        POIISD="09/2026", POIISDStatus="At Risk", EstIXCost="$7.4 M", IXCostStatus="Within Budget",
        SolarInverters="Power Electronics FS3430K / 30", MPT="150 MVA · Δ-Yg",
        StoragePCS="Sungrow PowerTitan / 50 / 4h", GenTieLength="0.8 miles",
        ReactivePowerComp="Cap Bank / 25 MVAr / 1", PowerPlantCtrl="Emerson",
        SafeHarbor="None", OtherAttributes="—",
        FutureExpansion="None", TransmissionNUs="Breaker upgrade / 08/2026",
        TopGatingItem="Recover 8-week MPT delivery slip",
    ),
    dict(
        ProjectID="P05", ReviewDate="07/01/2026", ProjectName="High Desert Hybrid",
        LeadInterconnection="O. Bennett", LeadHVEngineering="Y. Castillo", LeadSCADATelecom="Z. Haddad",
        LeadHVOM="Q. Larsen", LeadAssetMgmt="U. Mwangi",
        ProjectState="Nevada", TOUtility="NV Energy", ISORTO="WECC (non-RTO)",
        OTBEPAgreement="Yes (E&P 01/2024)", POIVoltage="230 kV", POIType="New SS",
        Solar="250 AC / 325 DC", Storage="125 MW / 500 MWh", PGIAGIA="03/2023", COD="10/2026",
        EPCHV="Mortenson", OMProvider="Pearce Renewables",
        POIISD="07/2026", POIISDStatus="Delayed", EstIXCost="$18.9 M", IXCostStatus="Over (+$2.4 M)",
        SolarInverters="TMEIC Ninja / 52", MPT="280 MVA · Δ-Yg",
        StoragePCS="BYD MC Cube / 125 / 4h", GenTieLength="4.6 miles",
        ReactivePowerComp="SVC / 60 MVAr / 1", PowerPlantCtrl="Trimark",
        SafeHarbor="MPT ~ 1 / 230kV-280MVA / FID 2022", OtherAttributes="NITS approved",
        FutureExpansion="Surplus IX 30 MW - DC", TransmissionNUs="None",
        TopGatingItem="SCADA FAT failure — retest 08/2026",
    ),
]

# (workstream, rowY, label, name, base timeline position 0-100)
MILESTONE_DEFS = [
    ("Commercial Analytics", 5, "C1", "Int. Basis / Curt", 5),
    ("Commercial Analytics", 5, "C2", "Ext. Basis / Curt", 16),
    ("Commercial Analytics", 5, "C3", "Study Refresh", 30),
    ("Commercial Analytics", 5, "C4", "3rd Party Study", 38),
    ("Commercial Analytics", 5, "C5", "Ext. Re-finance", 58),
    ("Commercial Analytics", 5, "C6", "Outage Monitoring (Basis/Curt.)", 74),
    ("Interconnection", 4, "I1", "IX App / POI", 5),
    ("Interconnection", 4, "I2", "IX Studies", 16),
    ("Interconnection", 4, "I3", "GIA Execution", 30),
    ("Interconnection", 4, "I4", "Dyn Settings (Analysis)", 38),
    ("Interconnection", 4, "I5", "Dyn Settings (Testing)", 58),
    ("Interconnection", 4, "I6", "As-Built Mods & LC Release", 74),
    ("HV Engineering", 3, "7", "MPT / HV Eqp Long-Lead", 2),
    ("HV Engineering", 3, "8", "E&P / OTB", 9),
    ("HV Engineering", 3, "9", "FacStudy / POI-HV Review", 28),
    ("HV Engineering", 3, "10", "Reactive & PQ (Studies)", 38),
    ("HV Engineering", 3, "11", "PEC", 50),
    ("HV Engineering", 3, "12", "V & Q Setting (Testing)", 58),
    ("HV Engineering", 3, "13", "As-Built SLD", 72),
    ("SCADA / Telecom", 2, "14", "GIA Req. Rev", 30),
    ("SCADA / Telecom", 2, "15", "PPA Req. Rev", 34),
    ("SCADA / Telecom", 2, "16", "Vendor Sel / Long-Lead Eqp", 38),
    ("SCADA / Telecom", 2, "17", "Points List Rev", 44),
    ("SCADA / Telecom", 2, "18", "TELCOM / Wless (Order/Install)", 48),
    ("SCADA / Telecom", 2, "19", "SCADA FAT", 54),
    ("SCADA / Telecom", 2, "20", "Primary Telco", 60),
    ("HV O&M", 1, "21", "Initial Design Rev. (ongoing)", 36),
    ("HV O&M", 1, "22", "HV Quality Rev. (ongoing)", 42),
    ("HV O&M", 1, "23", "MPT Del. / Inst.", 48),
    ("HV O&M", 1, "24", "Commissioning (CMX) Testing", 54),
    ("HV O&M", 1, "25", "Punch List Items", 58),
    ("HV O&M", 1, "26", "CC&C", 62),
]
WORKSTREAMS = ["Commercial Analytics", "Interconnection", "HV Engineering",
               "SCADA / Telecom", "HV O&M"]

# per project: progress cutoff (pos <= cutoff -> On Track), status overrides,
# and position shifts (to show circles being moved along the lifecycle)
PROJECT_TIMELINE = {
    "P00": dict(cutoff=-1, overrides={}, shifts={}),
    "P01": dict(cutoff=45, overrides={"11": "At Risk", "16": "At Risk"}, shifts={}),
    "P02": dict(cutoff=30, overrides={"I3": "At Risk", "C3": "At Risk"}, shifts={"I3": 4}),
    "P03": dict(cutoff=12, overrides={"9": "At Risk"}, shifts={}),
    "P04": dict(cutoff=40, overrides={"16": "Delayed", "11": "At Risk", "21": "At Risk"},
                shifts={"16": 3, "23": 3}),
    "P05": dict(cutoff=62, overrides={"19": "Delayed", "24": "Delayed",
                                      "20": "At Risk", "25": "At Risk"},
                shifts={"19": 9, "24": 12}),
}

KEY_ACTIONS = {
    "P00": [("[ Action item / decision needed ]", "[ Owner ]", "[ Due ]")] * 5,
    "P01": [
        ("Finalize PEC scope & schedule with Oncor", "K. Rivera", "07/15/2026"),
        ("Approve reactive power study addendum", "HV Eng", "07/22/2026"),
        ("Confirm SCADA points list freeze", "D. Okafor", "08/01/2026"),
        ("Sign off long-lead vendor QA plan", "A. Chen", "08/05/2026"),
        ("Refresh curtailment basis for PPA true-up", "Comm Analytics", "08/12/2026"),
    ],
    "P02": [
        ("IX cost overrun mitigation plan with APS", "S. Patel", "07/18/2026"),
        ("Re-baseline POI ISD (utility crew availability)", "S. Patel", "07/25/2026"),
        ("Decide OTB / E&P execution window", "C. Ferraro", "08/08/2026"),
        ("Complete basis / curtailment study refresh", "Comm Analytics", "08/15/2026"),
        ("Select PCS vendor — LOI", "R. Dominguez", "08/29/2026"),
    ],
    "P03": [
        ("Execute GIA with PJM / ComEd", "E. Kowalski", "08/30/2026"),
        ("Issue EPC RFP", "N. Adeyemi", "09/15/2026"),
        ("Convert MPT LOI to PO", "H. Yamada", "09/30/2026"),
        ("Kick off reactive & PQ studies", "H. Yamada", "10/15/2026"),
        ("Safe-harbor equipment allocation decision", "N. Adeyemi", "10/31/2026"),
    ],
    "P04": [
        ("Expedite plan for 8-week MPT delivery slip", "P. Lindqvist", "07/10/2026"),
        ("PEC readiness review with Georgia Power", "V. Osei", "07/24/2026"),
        ("Approve revised commissioning sequence", "I. Novak", "08/07/2026"),
        ("Close punch on gen-tie ROW restoration", "McCarthy", "08/21/2026"),
        ("Confirm breaker upgrade ISD with TO", "V. Osei", "08/28/2026"),
    ],
    "P05": [
        ("SCADA FAT retest — vendor corrective actions", "Z. Haddad", "08/01/2026"),
        ("Re-sequence CMX testing to protect COD", "Q. Larsen", "08/08/2026"),
        ("Negotiate IX cost true-up with NV Energy", "O. Bennett", "08/15/2026"),
        ("Punch list burn-down plan", "Q. Larsen", "09/01/2026"),
        ("CC&C documentation package prep", "M. Salinas", "09/15/2026"),
    ],
}


def milestone_rows():
    rows = []
    for proj in PROJECTS:
        pid = proj["ProjectID"]
        cfg = PROJECT_TIMELINE[pid]
        for sort, (ws, rowy, label, name, base) in enumerate(MILESTONE_DEFS, start=1):
            pos = base + cfg["shifts"].get(label, 0)
            if pid == "P00":
                status = "Not Started"
            else:
                status = "On Track" if pos <= cfg["cutoff"] else "Not Started"
                status = cfg["overrides"].get(label, status)
            rows.append([pid, ws, sort, label, name, pos, rowy, status])
    return rows


def write_csvs():
    os.makedirs(DATA_DIR, exist_ok=True)
    with open(os.path.join(DATA_DIR, "projects.csv"), "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(PROJECT_COLUMNS)
        for p in PROJECTS:
            w.writerow([p[c] for c in PROJECT_COLUMNS])
    with open(os.path.join(DATA_DIR, "milestones.csv"), "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(["ProjectID", "Workstream", "SortOrder", "Label",
                    "MilestoneName", "TimelinePos", "RowY", "Status"])
        w.writerows(milestone_rows())
    with open(os.path.join(DATA_DIR, "key_actions.csv"), "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(["ProjectID", "Seq", "ActionItem", "Owner", "Due"])
        for pid, items in KEY_ACTIONS.items():
            for i, (item, owner, due) in enumerate(items, start=1):
                w.writerow([pid, i, item, owner, due])


def write_xlsx():
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font
    except ImportError:
        print("openpyxl not available — skipping ProjectData.xlsx")
        return
    wb = Workbook()
    sheets = {
        "Projects": os.path.join(DATA_DIR, "projects.csv"),
        "Milestones": os.path.join(DATA_DIR, "milestones.csv"),
        "KeyActions": os.path.join(DATA_DIR, "key_actions.csv"),
    }
    first = True
    for name, path in sheets.items():
        ws = wb.active if first else wb.create_sheet()
        ws.title = name
        first = False
        with open(path, encoding="utf-8-sig") as f:
            for row in csv.reader(f):
                ws.append(row)
        for cell in ws[1]:
            cell.font = Font(bold=True)
        ws.freeze_panes = "A2"
    wb.save(os.path.join(DATA_DIR, "ProjectData.xlsx"))


# =====================================================================
# 2. SEMANTIC MODEL  (model.bim, TMSL)
# =====================================================================

def m_partition(csv_name, typed_cols):
    types = ", ".join('{"%s", %s}' % (c, t) for c, t in typed_cols)
    return (
        'let\n'
        '    Source = Csv.Document(File.Contents(DataFolder & "%s"), '
        '[Delimiter = ",", Encoding = 65001, QuoteStyle = QuoteStyle.Csv]),\n'
        '    #"Promoted Headers" = Table.PromoteHeaders(Source, [PromoteAllScalars = true]),\n'
        '    #"Changed Type" = Table.TransformColumnTypes(#"Promoted Headers", {%s})\n'
        'in\n'
        '    #"Changed Type"' % (csv_name, types)
    )


def build_model():
    proj_types = [(c, "type text") for c in PROJECT_COLUMNS]
    proj_cols = [dict(name=c, dataType="string", sourceColumn=c, summarizeBy="none")
                 for c in PROJECT_COLUMNS]
    proj_measures = [
        dict(name=m, expression='SELECTEDVALUE ( Projects[%s], "%s" )' % (col, ph))
        for col, m, _t, ph in FIELDS
    ]

    mile_types = [("ProjectID", "type text"), ("Workstream", "type text"),
                  ("SortOrder", "Int64.Type"), ("Label", "type text"),
                  ("MilestoneName", "type text"), ("TimelinePos", "type number"),
                  ("RowY", "Int64.Type"), ("Status", "type text")]
    mile_cols = [
        dict(name="ProjectID", dataType="string", sourceColumn="ProjectID", summarizeBy="none"),
        dict(name="Workstream", dataType="string", sourceColumn="Workstream", summarizeBy="none"),
        dict(name="SortOrder", dataType="int64", sourceColumn="SortOrder", summarizeBy="none"),
        dict(name="Label", dataType="string", sourceColumn="Label", summarizeBy="none",
             sortByColumn="SortOrder"),
        dict(name="MilestoneName", dataType="string", sourceColumn="MilestoneName", summarizeBy="none"),
        dict(name="TimelinePos", dataType="double", sourceColumn="TimelinePos", summarizeBy="none"),
        dict(name="RowY", dataType="int64", sourceColumn="RowY", summarizeBy="none"),
        dict(name="Status", dataType="string", sourceColumn="Status", summarizeBy="none"),
    ]
    mile_measures = [
        dict(name="Milestone Color", expression=(
            'SWITCH (\n'
            '    SELECTEDVALUE ( Milestones[Status] ),\n'
            '    "On Track", "#4E8542",\n'
            '    "At Risk", "#C4680C",\n'
            '    "Delayed", "#B3261E",\n'
            '    "Not Started", "#787E88",\n'
            '    "#787E88"\n'
            ')')),
    ]

    ka_types = [("ProjectID", "type text"), ("Seq", "Int64.Type"),
                ("ActionItem", "type text"), ("Owner", "type text"), ("Due", "type text")]
    ka_cols = [
        dict(name="ProjectID", dataType="string", sourceColumn="ProjectID", summarizeBy="none"),
        dict(name="Seq", dataType="int64", sourceColumn="Seq", summarizeBy="none"),
        dict(name="ActionItem", dataType="string", sourceColumn="ActionItem", summarizeBy="none"),
        dict(name="Owner", dataType="string", sourceColumn="Owner", summarizeBy="none"),
        dict(name="Due", dataType="string", sourceColumn="Due", summarizeBy="none"),
    ]

    def table(name, cols, csv_name, typed, measures=None):
        t = dict(
            name=name,
            columns=cols,
            partitions=[dict(name=name, mode="import",
                             source=dict(type="m", expression=m_partition(csv_name, typed)))],
        )
        if measures:
            t["measures"] = measures
        return t

    model = {
        "name": "Project Review",
        "compatibilityLevel": 1550,
        "model": {
            "culture": "en-US",
            "sourceQueryCulture": "en-US",
            "defaultPowerBIDataSourceVersion": "powerBI_V3",
            "dataAccessOptions": {"legacyRedirects": True, "returnErrorValuesAsNull": True},
            "expressions": [
                {
                    "name": "DataFolder",
                    "kind": "m",
                    "expression": '"C:\\ProjectReview\\data\\" meta [IsParameterQuery = true, Type = "Text", IsParameterQueryRequired = true]',
                    "annotations": [{"name": "PBI_ResultType", "value": "Text"}],
                }
            ],
            "tables": [
                table("Projects", proj_cols, "projects.csv", proj_types, proj_measures),
                table("Milestones", mile_cols, "milestones.csv", mile_types, mile_measures),
                table("KeyActions", ka_cols, "key_actions.csv", ka_types),
            ],
            "relationships": [
                {"name": "rel-milestones-projects", "fromTable": "Milestones",
                 "fromColumn": "ProjectID", "toTable": "Projects", "toColumn": "ProjectID"},
                {"name": "rel-keyactions-projects", "fromTable": "KeyActions",
                 "fromColumn": "ProjectID", "toTable": "Projects", "toColumn": "ProjectID"},
            ],
            "annotations": [
                {"name": "PBI_QueryOrder", "value": '["DataFolder","Projects","Milestones","KeyActions"]'},
                {"name": "__PBI_TimeIntelligenceEnabled", "value": "0"},
            ],
        },
    }
    return model


# =====================================================================
# 3. REPORT  (report.json, legacy PBIP layout)
# =====================================================================

_vid = [0]
_z = [0]


def next_name():
    _vid[0] += 1
    return "vc%05d" % _vid[0]


def next_z():
    _z[0] += 100
    return _z[0]


def lit(value):
    return {"expr": {"Literal": {"Value": value}}}


def slit(s):
    return lit("'%s'" % s.replace("'", "''"))


def dlit(n):
    return lit("%sD" % n)


def blit(b):
    return lit("true" if b else "false")


def solid(hex_color):
    return {"solid": {"color": lit("'%s'" % hex_color)}}


def solid_measure(entity, measure):
    return {"solid": {"color": {"expr": {"Measure": {
        "Expression": {"SourceRef": {"Entity": entity}}, "Property": measure}}}}}


def container(x, y, w, h, single_visual, filters=None):
    cfg = {
        "name": next_name(),
        "layouts": [{"id": 0, "position": {"x": x, "y": y, "z": next_z(),
                                           "width": w, "height": h}}],
        "singleVisual": single_visual,
    }
    vc = {"x": float(x), "y": float(y), "z": float(cfg["layouts"][0]["position"]["z"]),
          "width": float(w), "height": float(h), "config": json.dumps(cfg)}
    vc["filters"] = json.dumps(filters) if filters else "[]"
    return vc


def no_header(vc_objects=None):
    o = vc_objects or {}
    o["visualHeader"] = [{"properties": {"show": blit(False)}}]
    return o


def textbox(x, y, w, h, runs, bg=None, border=None, align="left"):
    paragraphs = [{"textRuns": [
        {"value": text, "textStyle": style} for text, style in runs
    ], "horizontalTextAlignment": align}]
    sv = {
        "visualType": "textbox",
        "drillFilterOtherVisuals": True,
        "objects": {"general": [{"properties": {"paragraphs": paragraphs}}]},
        "vcObjects": no_header(),
    }
    if bg:
        sv["vcObjects"]["background"] = [{"properties": {
            "show": blit(True), "color": solid(bg), "transparency": dlit(0)}}]
    if border:
        sv["vcObjects"]["border"] = [{"properties": {
            "show": blit(True), "color": solid(border), "radius": dlit(2)}}]
    return container(x, y, w, h, sv)


def style(size="9pt", color=INK, bold=False, italic=False):
    s = {"fontSize": size, "fontFamily": FONT, "color": color}
    if bold:
        s["fontWeight"] = "bold"
    if italic:
        s["fontStyle"] = "italic"
    return s


def section_header(x, y, w, text, h=20):
    return textbox(x, y, w, h, [(text, style("9pt", "#FFFFFF", bold=True))], bg=NAVY)


def card(x, y, w, h, measure, title, value_size=10, title_color=MUTED,
         value_color=INK, bg="#FFFFFF"):
    qref = "Projects.%s" % measure
    sv = {
        "visualType": "card",
        "projections": {"Values": [{"queryRef": qref}]},
        "prototypeQuery": {
            "Version": 2,
            "From": [{"Name": "p", "Entity": "Projects", "Type": 0}],
            "Select": [{"Measure": {"Expression": {"SourceRef": {"Source": "p"}},
                                    "Property": measure}, "Name": qref}],
        },
        "drillFilterOtherVisuals": True,
        "objects": {
            "labels": [{"properties": {
                "fontSize": dlit(value_size), "color": solid(value_color),
                "fontFamily": slit(FONT)}}],
            "categoryLabels": [{"properties": {"show": blit(False)}}],
            "wordWrap": [{"properties": {"show": blit(True)}}],
        },
        "vcObjects": no_header({
            "title": [{"properties": {
                "show": blit(True), "text": slit(title),
                "fontColor": solid(title_color), "fontSize": dlit(8),
                "fontFamily": slit(FONT)}}],
            "background": [{"properties": {
                "show": blit(True), "color": solid(bg), "transparency": dlit(0)}}],
            "border": [{"properties": {
                "show": blit(True), "color": solid(CARD_BORDER), "radius": dlit(3)}}],
        }),
    }
    return container(x, y, w, h, sv)


def col_select(alias, prop, entity_alias_map=None):
    return {"Column": {"Expression": {"SourceRef": {"Source": alias}}, "Property": prop}}


def filter_in(entity, prop, values, exclude=False, name="pf1"):
    cond = {"In": {
        "Expressions": [{"Column": {"Expression": {"SourceRef": {"Source": "f"}},
                                    "Property": prop}}],
        "Values": [[{"Literal": {"Value": "'%s'" % v}}] for v in values],
    }}
    if exclude:
        cond = {"Not": {"Expression": cond}}
    return {
        "name": name,
        "expression": {"Column": {"Expression": {"SourceRef": {"Entity": entity}},
                                  "Property": prop}},
        "filter": {"Version": 2,
                   "From": [{"Name": "f", "Entity": entity, "Type": 0}],
                   "Where": [{"Condition": cond}]},
        "type": "Categorical",
        "howCreated": 1,
    }


def table_visual(x, y, w, h, entity, alias, cols, display, order_by=None,
                 font_color_measure=None, filters=None, font_size=8,
                 header_bg="#EDEFF2"):
    selects, projections, col_props = [], [], {}
    for c, dn in zip(cols, display):
        qref = "%s.%s" % (entity, c)
        selects.append({"Column": {"Expression": {"SourceRef": {"Source": alias}},
                                   "Property": c}, "Name": qref})
        projections.append({"queryRef": qref})
        col_props[qref] = {"displayName": dn}
    proto = {"Version": 2, "From": [{"Name": alias, "Entity": entity, "Type": 0}],
             "Select": selects}
    if order_by:
        proto["OrderBy"] = [{"Direction": 1, "Expression": {
            "Column": {"Expression": {"SourceRef": {"Source": alias}},
                       "Property": order_by}}}]
    values_obj = [{"properties": {
        "fontSize": dlit(font_size), "fontFamily": slit(FONT),
        "wordWrap": blit(True)}}]
    if font_color_measure:
        for c in cols:
            qref = "%s.%s" % (entity, c)
            values_obj.append({
                "properties": {
                    "fontColorPrimary": solid_measure(entity, font_color_measure),
                    "fontColorSecondary": solid_measure(entity, font_color_measure),
                },
                "selector": {"metadata": qref},
            })
    sv = {
        "visualType": "tableEx",
        "projections": {"Values": projections},
        "prototypeQuery": proto,
        "columnProperties": col_props,
        "drillFilterOtherVisuals": True,
        "objects": {
            "grid": [{"properties": {
                "gridVertical": blit(False), "gridHorizontal": blit(True),
                "gridHorizontalColor": solid("#E4E6EA"), "rowPadding": dlit(2),
                "outlineColor": solid("#E4E6EA")}}],
            "columnHeaders": [{"properties": {
                "fontSize": dlit(font_size), "fontFamily": slit(FONT),
                "bold": blit(True), "fontColor": solid("#3B3F46"),
                "backColor": solid(header_bg), "wordWrap": blit(True)}}],
            "values": values_obj,
            "total": [{"properties": {"totals": blit(False)}}],
        },
        "vcObjects": no_header({
            "background": [{"properties": {
                "show": blit(True), "color": solid("#FFFFFF"), "transparency": dlit(0)}}],
            "border": [{"properties": {
                "show": blit(True), "color": solid(CARD_BORDER), "radius": dlit(3)}}],
        }),
    }
    return container(x, y, w, h, sv, filters=filters)


def timeline_scatter(x, y, w, h):
    sv = {
        "visualType": "scatterChart",
        "projections": {
            "Category": [{"queryRef": "Milestones.Label"}],
            "X": [{"queryRef": "Sum(Milestones.TimelinePos)"}],
            "Y": [{"queryRef": "Sum(Milestones.RowY)"}],
            "Tooltips": [{"queryRef": "Min(Milestones.Status)"},
                         {"queryRef": "Min(Milestones.MilestoneName)"}],
        },
        "prototypeQuery": {
            "Version": 2,
            "From": [{"Name": "m", "Entity": "Milestones", "Type": 0}],
            "Select": [
                {"Column": {"Expression": {"SourceRef": {"Source": "m"}},
                            "Property": "Label"}, "Name": "Milestones.Label"},
                {"Aggregation": {"Expression": {"Column": {
                    "Expression": {"SourceRef": {"Source": "m"}},
                    "Property": "TimelinePos"}}, "Function": 0},
                 "Name": "Sum(Milestones.TimelinePos)"},
                {"Aggregation": {"Expression": {"Column": {
                    "Expression": {"SourceRef": {"Source": "m"}},
                    "Property": "RowY"}}, "Function": 0},
                 "Name": "Sum(Milestones.RowY)"},
                {"Aggregation": {"Expression": {"Column": {
                    "Expression": {"SourceRef": {"Source": "m"}},
                    "Property": "Status"}}, "Function": 3},
                 "Name": "Min(Milestones.Status)"},
                {"Aggregation": {"Expression": {"Column": {
                    "Expression": {"SourceRef": {"Source": "m"}},
                    "Property": "MilestoneName"}}, "Function": 3},
                 "Name": "Min(Milestones.MilestoneName)"},
            ],
        },
        "drillFilterOtherVisuals": True,
        "objects": {
            "categoryAxis": [{"properties": {
                "show": blit(False), "start": lit("0L"), "end": lit("100L"),
                "showAxisTitle": blit(False), "gridlineShow": blit(False)}}],
            "valueAxis": [{"properties": {
                "show": blit(False), "start": lit("0L"), "end": lit("6L"),
                "showAxisTitle": blit(False), "gridlineShow": blit(True),
                "gridlineColor": solid("#EDEFF2")}}],
            "legend": [{"properties": {"show": blit(False)}}],
            "dataPoint": [{"properties": {
                "fill": solid_measure("Milestones", "Milestone Color")}}],
            "fillPoint": [{"properties": {"show": blit(True)}}],
            "markers": [{"properties": {"markerSize": dlit(7)}}],
            "bubbles": [{"properties": {"bubbleSize": dlit(30)}}],
            "categoryLabels": [{"properties": {
                "show": blit(True), "fontSize": dlit(8), "fontFamily": slit(FONT),
                "color": solid("#3B3F46")}}],
        },
        "vcObjects": no_header({
            "background": [{"properties": {
                "show": blit(True), "color": solid("#FFFFFF"), "transparency": dlit(0)}}],
        }),
    }
    return container(x, y, w, h, sv)


# ---------- page assembly ----------

# geometry shared by template + project pages
ROW_B_XS = [8, 170, 332, 494, 656, 818]
ROW_BC_W = 156

TL_X, TL_W = 100, 860           # scatter position/width
TL_Y, TL_H = 378, 244


def px_for_pos(pos):
    return TL_X + TL_W * pos / 100.0


def project_page(name, display_name, project_id, ordinal):
    vcs = []

    # --- top half: header cards ---
    vcs.append(card(8, 8, 170, 40, "v Review Date", "REVIEW DATE"))
    vcs.append(card(8, 52, 300, 52, "v Project Name", "PROJECT", value_size=16))

    lead_fields = [("v Lead IX", "INTERCONNECTION"), ("v Lead HV Eng", "HV ENGINEERING"),
                   ("v Lead SCADA", "SCADA / TELECOM"), ("v Lead HV OM", "HV O&M"),
                   ("v Lead Asset Mgmt", "ASSET MGMT")]
    for i, (m, t) in enumerate(lead_fields):
        vcs.append(card(320 + i * 136, 8, 130, 48, m, t))

    vcs.append(textbox(1004, 8, 268, 96,
                       [("POI SUBSTATION DESIGN\n", style("8pt", MUTED, bold=True)),
                        ("[ insert image / SLD ]", style("9pt", "#8A8F98", italic=True))],
                       bg="#FFFFFF", border=CARD_BORDER, align="center"))

    row_b = [("v Project State", "PROJECT STATE"), ("v TO Utility", "TO / UTILITY"),
             ("v ISO RTO", "ISO / RTO"), ("v OTB EP", "OTB / E&P AGREEMENT"),
             ("v POI Voltage", "POI VOLTAGE"), ("v POI Type", "POI TYPE")]
    for xpos, (m, t) in zip(ROW_B_XS, row_b):
        vcs.append(card(xpos, 110, ROW_BC_W, 44, m, t))

    row_c = [("v Solar", "SOLAR (MW AC / DC)"), ("v Storage", "STORAGE (MW / MWH)"),
             ("v PGIA GIA", "PGIA / GIA"), ("v COD", "COD"),
             ("v EPC HV", "EPC / HV"), ("v OM Provider", "O&M PROVIDER")]
    for xpos, (m, t) in zip(ROW_B_XS, row_c):
        vcs.append(card(xpos, 158, ROW_BC_W, 44, m, t))

    vcs.append(card(1004, 110, 268, 44, "v Top Gating Item", "TOP GATING ITEM"))
    vcs.append(card(1004, 158, 268, 44, "v Transmission NUs", "TRANSMISSION SYSTEM NUs"))

    vcs.append(section_header(8, 212, 300, "UTILITY PROGRESS"))
    vcs.append(section_header(316, 212, 560, "ENGINEERING  /  KEY DESIGN DETAILS"))
    vcs.append(section_header(884, 212, 388, "STRATEGIC  /  PROJECT ATTRIBUTES"))

    util = [("v POI ISD", "POI ISD"), ("v POI ISD Status", "POI ISD STATUS"),
            ("v Est IX Cost", "EST. IX COST"), ("v IX Cost Status", "IX COST STATUS")]
    for i, (m, t) in enumerate(util):
        vcs.append(card(8 + (i % 2) * 154, 236 + (i // 2) * 48, 146, 44, m, t))

    eng = [("v Solar Inverters", "SOLAR INVERTERS"), ("v MPT", "MPT"),
           ("v Storage PCS", "STORAGE / PCS"), ("v Gen Tie", "GEN-TIE LENGTH"),
           ("v Reactive Power", "REACTIVE POWER COMP"), ("v Power Plant Ctrl", "POWER PLANT CTRL")]
    for i, (m, t) in enumerate(eng):
        vcs.append(card(316 + (i % 3) * 188, 236 + (i // 3) * 48, 180, 44, m, t))

    strat = [("v Safe Harbor", "SAFE HARBOR"), ("v Other", "OTHER"),
             ("v Future Expansion", "FUTURE EXPANSION — NEW APPS"),
             ("v OTB EP", "OTB / E&P (DATES)")]
    for i, (m, t) in enumerate(strat):
        vcs.append(card(884 + (i % 2) * 196, 236 + (i // 2) * 48, 188, 44, m, t))

    # --- bottom half: milestone timeline ---
    vcs.append(textbox(8, 336, 380, 18,
                       [("CROSS-WORKSTREAM MILESTONE TIMELINE",
                         style("10pt", NAVY, bold=True))]))
    legend_runs = []
    for s, c in STATUS_COLORS.items():
        legend_runs.append(("  ● ", {"fontSize": "9pt", "fontFamily": FONT, "color": c,
                                     "fontWeight": "bold"}))
        legend_runs.append((s, style("8pt", "#3B3F46")))
    vcs.append(textbox(400, 336, 460, 18, legend_runs))

    phases = [("APP", 4), ("DEVELOPMENT", 15), ("GIA", 29), ("EPC", 37),
              ("CONSTRUCTION", 50), ("CMX", 63), ("COD", 73), ("COMMERCIAL OPS", 87)]
    for label, pos in phases:
        vcs.append(textbox(int(px_for_pos(pos)) - 52, 358, 104, 16,
                           [(label, style("8pt", NAVY, bold=True))], align="center"))

    for i, ws in enumerate(WORKSTREAMS):
        rowy = 5 - i
        center = TL_Y + TL_H * (6 - rowy) / 6.0
        vcs.append(textbox(8, int(center) - 14, 88, 30,
                           [(ws.upper(), style("7pt", "#FFFFFF", bold=True))], bg=NAVY))

    vcs.append(timeline_scatter(TL_X, TL_Y, TL_W, TL_H))

    vcs.append(section_header(976, 336, 296, "KEY ACTION ITEMS  /  NEXT STEPS"))
    vcs.append(table_visual(976, 360, 296, 262, "KeyActions", "k",
                            ["Seq", "ActionItem", "Owner", "Due"],
                            ["#", "Action item / decision needed", "Owner", "Due"],
                            order_by="Seq"))

    # --- bottom strip: per-workstream milestone status tables ---
    strip_xs = [8, 200, 392, 584, 776]
    strip_w = 186
    for xpos, ws in zip(strip_xs, WORKSTREAMS):
        yb = 630
        vcs.append(textbox(xpos, yb, strip_w, 16,
                           [(ws.upper(), style("7pt", "#FFFFFF", bold=True))], bg=NAVY))
        vf = [filter_in("Milestones", "Workstream", [ws], name="vf1")]
        vcs.append(table_visual(xpos, yb + 18, strip_w, 66, "Milestones", "m",
                                ["Label", "MilestoneName", "Status"],
                                ["#", "Milestone", "Status"],
                                order_by="Label",
                                font_color_measure="Milestone Color",
                                filters=vf, font_size=7))

    section = {
        "name": name,
        "displayName": display_name,
        "filters": json.dumps([filter_in("Projects", "ProjectID", [project_id])]),
        "ordinal": ordinal,
        "visualContainers": vcs,
        "config": json.dumps({}),
        "displayOption": 1,
        "width": float(PAGE_W),
        "height": float(PAGE_H),
    }
    return section


def summary_page(ordinal):
    vcs = []
    vcs.append(textbox(8, 10, 900, 34,
                       [("TRANSMISSION & GRID — PORTFOLIO SUMMARY",
                         style("16pt", NAVY, bold=True))]))
    vcs.append(textbox(8, 44, 900, 18,
                       [("Review date 07/01/2026  ·  5 active projects  ·  "
                         "milestone statuses are tracked on each project page",
                         style("9pt", MUTED, italic=True))]))

    vcs.append(section_header(8, 74, 1264, "FOUNDATIONAL PROJECT INFO"))
    vcs.append(table_visual(
        8, 98, 1264, 226, "Projects", "p",
        ["ProjectName", "ProjectState", "TOUtility", "ISORTO",
         "Solar", "Storage", "POIVoltage", "COD"],
        ["Project", "State", "TO / Utility", "ISO / RTO",
         "Solar (MW AC / DC)", "Storage (MW / MWh)", "POI Voltage", "COD"],
        order_by="ProjectName", font_size=10))

    vcs.append(section_header(8, 340, 1264, "TOP GATING ITEMS"))
    vcs.append(table_visual(
        8, 364, 1264, 200, "Projects", "p",
        ["ProjectName", "TopGatingItem", "COD"],
        ["Project", "Top gating item (placeholder — maintained in projects.csv)", "COD"],
        order_by="ProjectName", font_size=10))

    vcs.append(textbox(8, 580, 1264, 20,
                       [("Maintained from data/projects.csv — edit the CSV and refresh. "
                         "Circle statuses intentionally not shown here; see individual "
                         "project pages.", style("8pt", "#8A8F98", italic=True))]))

    return {
        "name": "SummaryPage",
        "displayName": "Portfolio Summary",
        "filters": json.dumps([filter_in("Projects", "ProjectID", ["P00"], exclude=True)]),
        "ordinal": ordinal,
        "visualContainers": vcs,
        "config": json.dumps({}),
        "displayOption": 1,
        "width": float(PAGE_W),
        "height": float(PAGE_H),
    }


def build_report():
    sections = [project_page("TemplatePage", "Template", "P00", 0)]
    for i, proj in enumerate(PROJECTS[1:], start=1):
        pid = proj["ProjectID"]
        sections.append(project_page(
            "Page%s" % pid, "%s · %s" % (pid, proj["ProjectName"]), pid, i))
    sections.append(summary_page(len(sections)))

    report_config = {
        "version": "5.43",
        "themeCollection": {"baseTheme": {"name": "CY24SU06", "version": "5.36", "type": 2}},
        "activeSectionIndex": 0,
        "defaultDrillFilterOtherVisuals": True,
        "settings": {"useNewFilterPaneExperience": True},
    }
    return {
        "config": json.dumps(report_config),
        "layoutOptimization": 0,
        "resourcePackages": [{"resourcePackage": {
            "disabled": False, "name": "SharedResources", "type": 2,
            "items": [{"name": "CY24SU06", "path": "BaseThemes/CY24SU06.json", "type": 202}],
        }}],
        "sections": sections,
        "publicCustomVisuals": [],
    }


# =====================================================================
# 4. SCAFFOLDING + WRITE
# =====================================================================

def write_json(path, obj):
    os.makedirs(os.path.dirname(path), exist_ok=True)
    with open(path, "w", encoding="utf-8") as f:
        json.dump(obj, f, indent=2, ensure_ascii=False)
        f.write("\n")


def main():
    write_csvs()
    write_xlsx()

    write_json(os.path.join(MODEL_DIR, "model.bim"), build_model())
    write_json(os.path.join(MODEL_DIR, "definition.pbism"), {
        "$schema": "https://developer.microsoft.com/json-schemas/fabric/item/semanticModel/definitionProperties/1.0.0/schema.json",
        "version": "4.0",
        "settings": {},
    })
    write_json(os.path.join(MODEL_DIR, ".platform"), {
        "$schema": "https://developer.microsoft.com/json-schemas/fabric/gitIntegration/platformProperties/2.0.0/schema.json",
        "metadata": {"type": "SemanticModel", "displayName": "Project Review"},
        "config": {"version": "2.0", "logicalId": "5c4c2c3e-9d1f-4a67-8e0a-1a2b3c4d5e01"},
    })

    write_json(os.path.join(REPORT_DIR, "report.json"), build_report())
    write_json(os.path.join(REPORT_DIR, "definition.pbir"), {
        "$schema": "https://developer.microsoft.com/json-schemas/fabric/item/report/definitionProperties/1.0.0/schema.json",
        "version": "1.0",
        "datasetReference": {"byPath": {"path": "../Project Review.SemanticModel"}},
    })
    write_json(os.path.join(REPORT_DIR, ".platform"), {
        "$schema": "https://developer.microsoft.com/json-schemas/fabric/gitIntegration/platformProperties/2.0.0/schema.json",
        "metadata": {"type": "Report", "displayName": "Project Review"},
        "config": {"version": "2.0", "logicalId": "5c4c2c3e-9d1f-4a67-8e0a-1a2b3c4d5e02"},
    })

    write_json(os.path.join(ROOT, "Project Review.pbip"), {
        "$schema": "https://developer.microsoft.com/json-schemas/fabric/pbip/pbipProperties/1.0.0/schema.json",
        "version": "1.0",
        "artifacts": [{"report": {"path": "Project Review.Report"}}],
        "settings": {"enableAutoRecovery": True},
    })

    # sanity: every emitted json parses, incl. embedded config/filter strings
    report = json.load(open(os.path.join(REPORT_DIR, "report.json"), encoding="utf-8"))
    json.loads(report["config"])
    n_visuals = 0
    for s in report["sections"]:
        json.loads(s["filters"])
        json.loads(s["config"])
        for vc in s["visualContainers"]:
            json.loads(vc["config"])
            json.loads(vc["filters"])
            n_visuals += 1
    json.load(open(os.path.join(MODEL_DIR, "model.bim"), encoding="utf-8"))
    rows = milestone_rows()
    print("OK: %d pages, %d visuals, %d milestone rows, %d projects"
          % (len(report["sections"]), n_visuals, len(rows), len(PROJECTS)))


if __name__ == "__main__":
    main()
